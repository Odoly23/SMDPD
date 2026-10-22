"""
Importa Database SEE-PLP (Excel) ba sistema: membru, departamentu, pozisaun iha struktura no rekursu intelektual.

	python manage.py import_see /dalan/Database_SEE-PLP_2026.xlsx --dry-run
	python manage.py import_see /dalan/Database_SEE-PLP_2026.xlsx [--aprova] [--orgaun SEE]

Aba "Form Responses 1" = fonte prinsipál (membru, pozisaun, departamentu, estrutura).
Aba "Sheet1" (formulariu Baze de Dadus Rekursu Intelektual) = esperensia, kontribuisaun, munisipiu orijin. Naran iha aba ne'e
la hanesan ho liña dadus nian, nune'e liña ida-ida *hatán* ba membru liu husi nu. telemovel (ka email, ka área+nivel úniku);
liña ne'ebé la bele hatán sei hatudu iha relatóriu no la importa.

Seguru atu la'o dala barak: membru hetan tuir nu. eleitoral ka naran; dadus ne'ebé iha ona la sobu-hakerek (deit hadia ne'ebé mamuk).
Membru foun hetan estatutu Pendente (uza --aprova atu Aprovadu). Dadus pesoál LA tama iha repositóriu: file Excel mantein deit iha ita-boot nia komputadór.
"""
import difflib
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.utils import timezone
from custom.models import Municipality, AdministrativePost, Village, Aldeia, Departamentu, Organizasaun
from membro.models import Membro, Alokasaun, VICE_BARAK, POZISAUN_POR_NIVEL
from membro.utils import log_historia

EDUKASAUN = {'sd': 'SD', 'smp': 'SMP', 'sma': 'SMA', 'd3': 'D3', 's1': 'S1', 's2': 'S2', 's3': 'S3', 'kandidatus1': 'Kandidatu S1'}
ALIAS_MUN = {'oecusi': 'oecusse', 'oecusse': 'oecusse', 'manatutu': 'manatuto', 'ambeno': 'oecusse'}
ROMANU = {'i': 1, 'ii': 2, 'iii': 3, 'iv': 4, 'v': 5, 'vi': 6, 'vii': 7, 'viii': 8}
PREFIXU_DEP = ('asuntu', 'asunto', 'divisaun', 'divisao', 'seitor', 'sektor', 'apoiu')
HAKARAK_SHEET1 = 'Deskrisaun konaba Esperensia Servisu'


def norma(s):
	t = unicodedata.normalize('NFKD', str(s if s is not None else '')).encode('ascii', 'ignore').decode().lower()
	return re.sub(r'[^a-z0-9]', '', t)


def limpa(s):
	return re.sub(r'\s+', ' ', str(s if s is not None else '').replace('\xa0', ' ')).strip()


def kuaa(s):
	"""Naran ba lookup munisipiu/postu/suku/aldeia (la sensitivu ba asentu, espasu no alias)."""
	n = norma(s)
	n = re.sub(r'^(municipio|munisipiu|posto|postu)', '', n) if n.startswith(('municipio', 'munisipiu')) else n
	if n.startswith('oecusse') or n.startswith('oecusi'):
		return 'oecusse'
	return ALIAS_MUN.get(n, n)


def chave_dep(nome):
	n = limpa(nome).lower()
	n = unicodedata.normalize('NFKD', n).encode('ascii', 'ignore').decode()
	for p in PREFIXU_DEP:
		if n.startswith(p + ' '):
			n = n[len(p) + 1:]
			break
	n = re.sub(r'[^a-z0-9]', '', n)
	return n.replace('ss', 's').replace('sasaun', 'saun')


def pozisaun_ba(raw):
	"""(tipu, nu, ekstra) husi testu pozisaun: tipu = Presidente|Koordenador|Vice|Membru ka None."""
	t = limpa(raw)
	if not t:
		return None, None, ''
	n = unicodedata.normalize('NFKD', t.lower()).encode('ascii', 'ignore').decode()
	if 'presidente' in n:
		return 'Presidente', None, ''
	if 'vice' in n or 'vise' in n:
		nu = None
		m = re.search(r'\b(\d)\b', n)
		if m:
			nu = int(m.group(1))
		else:
			m = re.search(r'\b(viii|vii|vi|iv|iii|ii|i|v)\b', n)
			if m:
				nu = ROMANU[m.group(1)]
		ekstra = ''
		m = re.search(r'(?:asuntu|asunto)\s+(.+)$', t, re.I)
		if m:
			ekstra = limpa(m.group(1))
		return 'Vice', nu, ekstra
	if 'koord' in n or 'kord' in n:
		return 'Koordenador', None, ''
	return 'Membru', None, ''


class Command(BaseCommand):
	help = "Importa Database SEE-PLP (Excel): membru, departamentu, pozisaun no rekursu intelektual."

	def add_arguments(self, parser):
		parser.add_argument('xlsx', help='File Excel (Database_SEE-PLP_2026.xlsx)')
		parser.add_argument('--orgaun', default='SEE', help='Kódigu orgaun/movimentu (default: SEE)')
		parser.add_argument('--aprova', action='store_true', help='Membru foun hetan estatutu Aprovadu (defeitu: Pendente)')
		parser.add_argument('--dry-run', action='store_true', help='Haree rezultadu deit, la rai buat ida')

	# ───────────────────────── leitura ─────────────────────────
	def le_aba(self, wb, naran, marka):
		if naran not in wb.sheetnames:
			return []
		linhas = [list(r) for r in wb[naran].iter_rows(values_only=True)]
		for i, r in enumerate(linhas):
			if marka in [limpa(c) for c in r]:
				kab = [limpa(c) for c in r]
				return [dict(zip(kab, l)) for l in linhas[i + 1:] if any(c not in (None, '') for c in l)]
		return []

	def handle(self, *args, **opt):
		try:
			from openpyxl import load_workbook
		except ImportError:
			raise CommandError('Instala openpyxl: pip install openpyxl')
		try:
			wb = load_workbook(opt['xlsx'], read_only=True, data_only=True)
		except Exception as e:
			raise CommandError(f'La bele loke file: {e}')
		resp = [r for r in self.le_aba(wb, 'Form Responses 1', 'Naran Kompletu') if limpa(r.get('Naran Kompletu'))]
		if not resp:
			raise CommandError('Aba "Form Responses 1" ho koluna "Naran Kompletu" la hetan.')
		sheet1 = [r for r in self.le_aba(wb, 'Sheet1', HAKARAK_SHEET1) if limpa(r.get('Naran\n(tuir ordem alfabetiku)') or r.get('Naran (tuir ordem alfabetiku)') or r.get('Naran'))]
		self.avizu = defaultdict(list)
		with transaction.atomic():
			res = self.importa(resp, sheet1, opt)
			if opt['dry_run']:
				transaction.set_rollback(True)
		self.relatoriu(res, opt['dry_run'])

	# ───────────────────────── importa ─────────────────────────
	def importa(self, resp, sheet1, opt):
		org = Organizasaun.objects.filter(code__iexact=opt['orgaun']).first()
		if not org:
			if opt['orgaun'].upper() != 'SEE':
				raise CommandError(f'Orgaun "{opt["orgaun"]}" la iha. Kria uluk iha Data Master.')
			org = Organizasaun.objects.create(code='SEE', name='Sentru Estudu Estratejiku')
		if org.code.upper() == 'SEE':
			if not org.titulu:
				org.titulu = 'Orgaun Konsultivu | Orgão Consultivo PLP'
			if not org.sede:
				org.sede = 'Sede Kakaulidun, Bairo Pite, Dom Aleixo, Dili, Timor-Leste'
			org.save()
		self.org = org
		self.mun = {kuaa(m.name): m for m in Municipality.objects.all()}
		self.res = Counter()
		self.dep_foun = []
		self.deps = {chave_dep(d.name): d for d in Departamentu.objects.filter(organizasaun=org)}
		self.ada_org = {norma(m.nome): m for m in Membro.objects.filter(organizasaun=org)}
		self.por_nu = {m.nu_eleitoral: m for m in Membro.objects.exclude(nu_eleitoral__isnull=True)}
		self.aprova = opt['aprova']

		pesoa = {}          # naran normalizadu -> {'dados': {...}, 'poz': [...]}
		for r in resp:
			nome = limpa(r.get('Naran Kompletu'))
			k = norma(nome)
			if k not in pesoa:
				mirip = next((x for x in pesoa if difflib.SequenceMatcher(None, x, k).ratio() >= 0.95), None)
				if mirip:
					self.avizu['Naran hanesan ho lia-fuan ki\'ik (hamutuk)'].append(f"{pesoa[mirip]['nome']} ≈ {nome} (liña {r.get('No.')})")
					k = mirip
			p = pesoa.setdefault(k, {'nome': nome, 'dados': {}, 'poz': [], 'linhas': []})
			p['linhas'].append(r.get('No.'))
			self.kumpleta(p['dados'], self.dadus_membru(r))
			if r.get('Pozisaun') or r.get('Departamentu') or r.get('Estrutura'):
				p['poz'].append((r.get('Pozisaun'), r.get('Departamentu'), r.get('Estrutura')))
		self.res['liña Excel'] = len(resp)
		self.res['pesoa úniku'] = len(pesoa)
		for k, p in pesoa.items():
			if len(p['linhas']) > 1:
				self.avizu['Naran repete (hamutuk)'].append(f"{p['nome']} (liña {', '.join(str(x) for x in p['linhas'])})")

		self.sheet1(pesoa, sheet1)
		self.membros = {}
		for k, p in pesoa.items():
			self.membros[k] = self.rai_membru(p)
		self.unidade = {}   # (nivel, dep, mun, pozisaun) -> membru
		for k, p in pesoa.items():
			for poz in p['poz']:
				self.rai_pozisaun(self.membros[k], p['nome'], *poz)
		return self.res

	# ───────────────────────── membru ─────────────────────────
	def dadus_membru(self, r):
		d = {}
		sex = limpa(r.get('Sexu')).capitalize()
		d['sex'] = sex if sex in ('Mane', 'Feto') else 'Mane'
		if sex not in ('Mane', 'Feto'):
			self.avizu['Sexu la loos (uza Mane)'].append(limpa(r.get('Naran Kompletu')))
		ele = limpa(r.get('Nu. Eleitoral')).upper()
		d['nu_eleitoral'] = ele or None
		dob = r.get('Data Moris')
		d['date_of_birth'] = dob.date() if isinstance(dob, datetime) else dob if isinstance(dob, date) else None
		tel = re.sub(r'[^0-9+]', '', limpa(r.get('Numero Telemovel / WhatsApp')))
		d['telefone'] = tel if 7 <= len(tel) <= 20 else None
		mail = limpa(r.get('Email')).lower()
		d['email'] = mail if re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', mail) else None
		d['nivel_edukasaun'] = EDUKASAUN.get(norma(r.get('Nivel Edukusaun ') or r.get('Nivel Edukusaun')))
		d['area_estudu'] = limpa(r.get('Area Estudu/Expert')) or None
		d['address'] = limpa(r.get('Hela Fatin Agora'))[:150] or None
		d['servisu_fatin'] = limpa(r.get('Servisu Fatin Agora'))[:150] or None
		geo = self.geografia(r)
		d.update(geo)
		return {k: v for k, v in d.items() if v not in (None, '')}

	def geografia(self, r):
		nome = limpa(r.get('Naran Kompletu'))
		out = {}
		m = limpa(r.get('Municipiu'))
		if not m:
			return out
		mun = self.mun.get(kuaa(m))
		if not mun:
			self.avizu['Munisipiu la hetan'].append(f'{nome}: {m}')
			return out
		out['municipality'] = mun
		p = limpa(r.get('Posto Administrativu'))
		if p:
			post = next((x for x in AdministrativePost.objects.filter(municipality=mun) if norma(x.name) == norma(p)), None)
			if post:
				out['post'] = post
				s = limpa(r.get('Suku'))
				if s:
					suku = next((x for x in Village.objects.filter(administrativepost=post) if norma(x.name) == norma(s)), None)
					if suku:
						out['village'] = suku
						a = limpa(r.get('Aldeia'))
						if a:
							ald = next((x for x in Aldeia.objects.filter(village=suku) if norma(x.name) == norma(a)), None)
							if ald:
								out['aldeia'] = ald
							else:
								self.avizu['Aldeia la hetan'].append(f'{nome}: {a} ({s})')
					else:
						self.avizu['Suku la hetan'].append(f'{nome}: {s} ({p})')
			else:
				self.avizu['Postu la hetan'].append(f'{nome}: {p} ({mun.name})')
		return out

	@staticmethod
	def kumpleta(alvu, fonte):
		for k, v in fonte.items():
			if alvu.get(k) in (None, ''):
				alvu[k] = v

	def sheet1(self, pesoa, linhas):
		"""Hatán liña formulariu ba membru: nu. telemovel > email > (área + nivel úniku)."""
		if not linhas:
			return
		por_tel = defaultdict(list)
		por_mail = defaultdict(list)
		por_an = defaultdict(list)
		for p in pesoa.values():
			d = p['dados']
			if d.get('telefone'):
				por_tel[d['telefone']].append(p)
			if d.get('email'):
				por_mail[d['email']].append(p)
			if d.get('area_estudu') and d.get('nivel_edukasaun'):
				por_an[(norma(d['area_estudu']), d['nivel_edukasaun'])].append(p)
		ok = Counter()
		for r in linhas:
			tel = re.sub(r'[^0-9+]', '', limpa(r.get('Nu. Kontaktu')))
			mail = limpa(r.get('Email')).lower()
			area = limpa(r.get('Area Estudu'))
			nivel = EDUKASAUN.get(norma(r.get('Nivel Estudu / Gelar')))
			extra = {'esperensia': limpa(r.get(HAKARAK_SHEET1)) or None, 'kontribuisaun': limpa(r.get('Kontribuisaun ba Programa PLP')) or None}
			orijin = limpa(r.get('Munisipiu Orijin'))
			rezid = limpa(r.get('Munisipiu Rezidensia'))
			if orijin and self.mun.get(kuaa(orijin)):
				extra['municipality_orijin'] = self.mun[kuaa(orijin)]
			if rezid and self.mun.get(kuaa(rezid)):
				extra['municipality'] = self.mun[kuaa(rezid)]
			if area:
				extra['area_estudu'] = area
			if nivel:
				extra['nivel_edukasaun'] = nivel
			extra['servisu_fatin'] = None
			if not any(v for v in extra.values()):
				ok['mamuk'] += 1
				continue
			alvu = None
			if tel and len(por_tel.get(tel, [])) == 1:
				alvu, ok_ba = por_tel[tel][0], 'telemovel'
			elif mail and len(por_mail.get(mail, [])) == 1:
				alvu, ok_ba = por_mail[mail][0], 'email'
			elif area and nivel and len(por_an.get((norma(area), nivel), [])) == 1:
				alvu, ok_ba = por_an[(norma(area), nivel)][0], 'área+nivel'
			if not alvu:
				ok['la hatán'] += 1
				self.avizu['Liña formulariu la bele hatán ba membru (la importa)'].append(
					f"No {r.get('No')}: {limpa(r.get('Area Estudu')) or '-'} / {limpa(r.get('Nu. Kontaktu')) or '-'}")
				continue
			ada = alvu['dados'].get('area_estudu')
			if ada and area and norma(ada) != norma(area):
				ok['konflitu'] += 1
				self.avizu['Liña formulariu ho área estudu la hanesan (la importa)'].append(f"No {r.get('No')}: '{area}' ≠ '{ada}' ({alvu['nome']})")
				continue
			self.kumpleta(alvu['dados'], {k: v for k, v in extra.items() if v})
			ok['hatán liu ' + ok_ba] += 1
		for k, v in ok.items():
			self.res['formulariu: ' + k] = v

	def rai_membru(self, p):
		d = dict(p['dados'])
		ele = d.get('nu_eleitoral')
		m = self.por_nu.get(ele) if ele else None
		if m is None:
			m = self.ada_org.get(norma(p['nome']))
		if m is None:
			if ele and ele in self.por_nu:
				d.pop('nu_eleitoral', None)
			m = Membro(nome=p['nome'], organizasaun=self.org, status='Ativu', datetime=timezone.now(),
				aprovasaun='Aprovadu' if self.aprova else 'Pendente')
			if self.aprova:
				m.data_aprovasaun = timezone.now()
			for k, v in d.items():
				setattr(m, k, v)
			m.save()
			log_historia(None, m, 'Importa', f'Importa Database SEE: {m}')
			self.res['membru foun'] += 1
		else:
			mudou = False
			if m.organizasaun_id is None:
				m.organizasaun = self.org
				mudou = True
			for k, v in d.items():
				if getattr(m, k) in (None, ''):
					if k == 'nu_eleitoral' and Membro.objects.filter(nu_eleitoral=v).exclude(pk=m.pk).exists():
						continue
					setattr(m, k, v)
					mudou = True
			if mudou:
				m.save()
				self.res['membru atualiza'] += 1
			else:
				self.res['membru la muda'] += 1
		if m.nu_eleitoral:
			self.por_nu[m.nu_eleitoral] = m
		self.ada_org[norma(m.nome)] = m
		return m

	# ───────────────────────── pozisaun ─────────────────────────
	def departamentu(self, naran):
		k = chave_dep(naran)
		d = self.deps.get(k)
		if d is None:
			n = Departamentu.objects.filter(organizasaun=self.org).count() + 1
			d = Departamentu.objects.create(organizasaun=self.org, code=f'D{n:02d}', name=limpa(naran)[:100])
			self.deps[k] = d
			self.dep_foun.append(d.name)
		return d

	def rai_pozisaun(self, membru, nome, poz_raw, dep_raw, est_raw):
		tipu, nu, ekstra = pozisaun_ba(poz_raw)
		dep = limpa(dep_raw)
		dn = norma(dep)
		est = limpa(est_raw)
		if dn in ('', '-', 'jeral') or 'vice' in dn:
			if dep and 'vice' in dn:
				self.avizu['Departamentu la klaru (hela deit membru)'].append(f'{nome}: {dep}')
				tipu = 'Membru' if tipu == 'Koordenador' else tipu
			dep = ''
		if not dep and ekstra:
			dep = ekstra       # ez. "Vice Koordenador Asuntu Kurikulo"
		municipality = None
		if est.lower().startswith(('municipio', 'munisipiu')):
			municipality = self.mun.get(kuaa(est))
			if not municipality:
				self.avizu['Munisipiu estrutura la hetan'].append(f'{nome}: {est}')
				return
		if tipu is None:
			if not (municipality or dep):
				return
			tipu = 'Membru'
		if municipality:
			dep = ''
		if dep:
			nivel, depobj = 'Departamentu', self.departamentu(dep)
			pozisaun_nama = {'Presidente': 'Presidente Departamentu'}.get(tipu, tipu)
		else:
			nivel, depobj = 'Organizasaun', None
			pozisaun_nama = tipu
		if tipu == 'Membru' and nivel == 'Organizasaun' and not municipality:
			return      # membru orgaun: ona rejistu liu husi orgaun/movimentu
		if tipu == 'Vice':
			pozisaun_nama = self.vice_livre(nivel, depobj, municipality, nu, membru, nome)
			if not pozisaun_nama:
				return
		chave = (nivel, depobj.pk if depobj else None, municipality.pk if municipality else None, pozisaun_nama)
		if pozisaun_nama != 'Membru':
			atual = self.unidade.get(chave)
			if atual and atual.pk == membru.pk:
				return                                          # pozisaun hanesan ba membru hanesan
			if atual:
				self.avizu['Pozisaun okupadu ona (sai Membru)'].append(f'{nome}: {pozisaun_nama} {depobj or self.org}{" " + municipality.name if municipality else ""} (okupa: {atual.nome})')
				pozisaun_nama = 'Membru'
				chave = (nivel, depobj.pk if depobj else None, municipality.pk if municipality else None, 'Membru')
			else:
				self.unidade[chave] = membru
		if pozisaun_nama == 'Membru' and self.unidade.get(chave + (membru.pk,)):
			return
		kriteriu = {'membro': membru, 'nivel': nivel, 'pozisaun': pozisaun_nama, 'departamentu': depobj,
			'organizasaun': None if nivel == 'Departamentu' else self.org, 'municipality': municipality}
		_, foun = Alokasaun.objects.get_or_create(**kriteriu, defaults={'data_hahu': date.today(), 'datetime': timezone.now()})
		self.unidade[chave + (membru.pk,)] = membru if pozisaun_nama == 'Membru' else None
		self.res['pozisaun foun' if foun else 'pozisaun iha ona'] += 1

	def vice_livre(self, nivel, dep, municipality, nu, membru, nome):
		"""Vice nu. 1-8: uza nu. husi Excel; se okupadu/la iha nu., hili ne'ebé livre."""
		def chave(v):
			return (nivel, dep.pk if dep else None, municipality.pk if municipality else None, v)
		opsaun = [f'Vice {nu}'] if nu and 1 <= nu <= 8 else []
		opsaun += [v for v in VICE_BARAK if v not in opsaun]
		for v in opsaun:
			atual = self.unidade.get(chave(v))
			if atual and atual.pk == membru.pk:
				return v
			if not atual:
				if nu and v != f'Vice {nu}':
					self.avizu['Vice nu. okupadu (hili seluk)'].append(f'{nome}: Vice {nu} -> {v}')
				return v
		self.avizu['Vice hotu okupadu'].append(nome)
		return None

	# ───────────────────────── relatóriu ─────────────────────────
	def relatoriu(self, res, dry):
		out = self.stdout
		out.write(self.style.SUCCESS(('[DRY-RUN — la rai buat ida] ' if dry else '') + 'Rezultadu:'))
		for k, v in res.items():
			out.write(f'  · {k}: {v}')
		if self.dep_foun:
			out.write(self.style.SUCCESS(f'  · departamentu foun ({len(self.dep_foun)}): ' + '; '.join(self.dep_foun)))
		for titulu, lista in self.avizu.items():
			out.write(self.style.WARNING(f'\nAVIZU — {titulu} ({len(lista)}):'))
			for x in lista[:12]:
				out.write(f'   - {x}')
			if len(lista) > 12:
				out.write(f'   … no {len(lista) - 12} tan')
