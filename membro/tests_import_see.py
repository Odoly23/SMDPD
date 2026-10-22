import datetime
import os
import tempfile
from io import StringIO
from django.core.management import call_command
from django.core.management.base import CommandError
from django.test import TestCase
from openpyxl import Workbook
from custom.models import Municipality, AdministrativePost, Village, Aldeia, Departamentu, Organizasaun
from membro.models import Membro, Alokasaun

H1 = [None, 'No.', 'Nu. Eleitoral', 'Naran Kompletu', 'Data Moris', 'Sexu', 'Email', 'Numero Telemovel / WhatsApp', 'Pozisaun', 'Departamentu',
	'Nivel Edukusaun ', 'Municipiu', 'Posto Administrativu', 'Suku', 'Aldeia', 'Area Estudu/Expert', 'Hela Fatin Agora', 'Servisu Fatin Agora', 'Estrutura']
H2 = [None, None, 'No', 'Naran\n(tuir ordem alfabetiku)', 'Area Estudu', 'Nivel Estudu / Gelar', 'Deskrisaun konaba Esperensia Servisu',
	'Kontribuisaun ba Programa PLP', 'Munisipiu Orijin', 'Munisipiu Rezidensia', 'Nu. Kontaktu', 'Email', None]

# naran/dadus FIKTISIU (laiha dadus pesoál loos)
LINHAS = [
	# nu, eleitoral, naran, dob, sex, email, tel, pozisaun, dep, nivel, mun, postu, suku, aldeia, area, hela, servisu, estrutura
	(1, 'EL-001', 'Aaa Presidente Uji', datetime.datetime(1980, 1, 2), 'Mane', 'aaa@teste.tl', 77000001, 'Presidente', 'Jeral', 'S2', 'Dili', 'Dom Aleixo', 'Comoro', 'Aldeia 1', 'Ekonomia', 'Comoro', 'Kantor X', 'Nasional'),
	(2, None, 'Bbb Koord Finansas', None, 'Feto', None, 77000002, 'Koordenador', 'Asuntu Finansas', 'S1', None, None, None, None, 'Kontabilidade', None, None, 'Nasional'),
	(3, None, 'Ccc Vice Finansas', None, 'Mane', None, 77000003, 'Vice Koordenador I', 'Asunto Finansas', None, None, None, None, None, None, None, None, 'Nasional'),
	(4, None, 'Ddd Vice Finansas Dua', None, 'Mane', None, None, 'Vice Koordenador', 'Finansas', None, None, None, None, None, None, None, None, 'Nasional'),
	(5, None, 'Eee Koord Finansas Seluk', None, 'Mane', None, None, 'Koordenador', 'Asuntu Finansas', None, None, None, None, None, None, None, None, 'Nasional'),
	(6, None, 'Fff Koord Jeral Oecusi', None, 'Mane', None, None, 'Koordenador', 'Jeral', None, None, None, None, None, None, None, None, 'Municipio Oecusi'),
	(7, None, 'Ggg Membru Nasionál', None, 'Feto', None, None, 'Membru', None, None, None, None, None, None, None, None, None, 'Nasional'),
	(8, None, 'Hhh Sein Pozisaun', None, 'Mane', None, 77000008, None, None, 'SMA', 'Dili', 'Postu La Iha', None, None, 'Ingles', None, None, None),
	(9, None, 'Iii Vice Oito', None, 'Mane', None, None, 'Vise 8', 'Asunto Social', None, None, None, None, None, None, None, None, 'Nasional'),
]
SHEET1 = [  # data hatán liu telemovel; naran iha kolona la hanesan (hanesan ficheiru loos)
	(1, 'Naran Seluk', 'Ekonomia', 'S2', 'Aaa pengalaman ONG', 'Hadia ekonomia', 'Baucau', 'Dili', 77000001, None),
	(2, 'Naran Seluk Dua', 'Teknolojia', 'S1', 'Programador', None, 'Lautem', None, 99999999, None),   # la hatán
]


def buat_xlsx(path, linhas=LINHAS, sheet1=SHEET1):
	wb = Workbook()
	ws = wb.active
	ws.title = 'Form Responses 1'
	ws.append([None, 'BASE DE DADUS ...'])
	ws.append(H1)
	for l in linhas:
		ws.append([None, l[0], l[1], l[2], l[3], l[4], l[5], l[6], l[7], l[8], l[9], l[10], l[11], l[12], l[13], l[14], l[15], l[16], l[17]])
	s2 = wb.create_sheet('Sheet1')
	s2.append([None, None, 'Orgaun Konsultivu'])
	s2.append([])
	s2.append(H2)
	for r in sheet1:
		s2.append([None, None, r[0], r[1], r[2], r[3], r[4], r[5], r[6], r[7], r[8], r[9]])
	wb.save(path)


class ImportSeeTests(TestCase):
	@classmethod
	def setUpTestData(cls):
		cls.dili = Municipality.objects.create(code='DIL', name='Dili')
		cls.oe = Municipality.objects.create(code='OEC', name='Oé-Cusse Ambeno (RAEOA)')
		cls.baucau = Municipality.objects.create(code='BAU', name='Baucau')
		cls.post = AdministrativePost.objects.create(municipality=cls.dili, name='Dom Aleixo')
		cls.vil = Village.objects.create(administrativepost=cls.post, name='Comoro')
		cls.ald = Aldeia.objects.create(village=cls.vil, name='Aldeia 1')
		cls.org = Organizasaun.objects.create(code='SEE', name='Sentru Estudu Estratejiku')

	def setUp(self):
		fd, self.path = tempfile.mkstemp(suffix='.xlsx')
		os.close(fd)
		buat_xlsx(self.path)
		self.addCleanup(os.remove, self.path)

	def jalan(self, *args):
		out = StringIO()
		call_command('import_see', self.path, *args, stdout=out)
		return out.getvalue()

	def test_membru_no_dadus(self):
		self.jalan()
		self.assertEqual(Membro.objects.count(), 9)
		a = Membro.objects.get(nome='Aaa Presidente Uji')
		self.assertEqual((a.nu_eleitoral, str(a.date_of_birth), a.sex, a.telefone, a.email, a.nivel_edukasaun, a.area_estudu),
			('EL-001', '1980-01-02', 'Mane', '77000001', 'aaa@teste.tl', 'S2', 'Ekonomia'))
		self.assertEqual((a.municipality, a.post, a.village, a.aldeia, a.organizasaun, a.aprovasaun, a.status),
			(self.dili, self.post, self.vil, self.ald, self.org, 'Pendente', 'Ativu'))
		self.assertEqual((a.address, a.servisu_fatin), ('Comoro', 'Kantor X'))
		b = Membro.objects.get(nome='Bbb Koord Finansas')
		self.assertIsNone(b.nu_eleitoral)
		self.assertIsNone(b.date_of_birth)
		self.assertFalse(b.kompletu())

	def test_orgaun_titulu_no_sede(self):
		self.jalan()
		self.org.refresh_from_db()
		self.assertIn('Orgaun Konsultivu', self.org.titulu)
		self.assertIn('Bairo Pite', self.org.sede)

	def test_pozisaun_organizasaun_no_departamentu(self):
		self.jalan()
		pres = Alokasaun.objects.get(membro__nome='Aaa Presidente Uji')
		self.assertEqual((pres.nivel, pres.pozisaun, pres.organizasaun, pres.departamentu, pres.municipality_id), ('Organizasaun', 'Presidente', self.org, None, None))
		fin = Departamentu.objects.get(organizasaun=self.org, name__icontains='Finansas')
		self.assertEqual(Departamentu.objects.filter(organizasaun=self.org, name__icontains='Finansas').count(), 1)  # varians hamutuk
		k = Alokasaun.objects.get(membro__nome='Bbb Koord Finansas')
		self.assertEqual((k.nivel, k.pozisaun, k.departamentu), ('Departamentu', 'Koordenador', fin))
		v1 = Alokasaun.objects.get(membro__nome='Ccc Vice Finansas')
		v2 = Alokasaun.objects.get(membro__nome='Ddd Vice Finansas Dua')
		self.assertEqual((v1.pozisaun, v2.pozisaun), ('Vice 1', 'Vice 2'))
		v8 = Alokasaun.objects.get(membro__nome='Iii Vice Oito')
		self.assertEqual((v8.pozisaun, v8.departamentu.name), ('Vice 8', 'Asunto Social'))

	def test_pozisaun_okupadu_sai_membru(self):
		out = self.jalan()
		e = Alokasaun.objects.get(membro__nome='Eee Koord Finansas Seluk')
		self.assertEqual(e.pozisaun, 'Membru')
		self.assertIn('okupadu', out.lower())

	def test_estrutura_munisipiu(self):
		self.jalan()
		a = Alokasaun.objects.get(membro__nome='Fff Koord Jeral Oecusi')
		self.assertEqual((a.nivel, a.pozisaun, a.municipality, a.organizasaun), ('Organizasaun', 'Koordenador', self.oe, self.org))
		self.assertIn('Oé-Cusse', str(a.unit()))

	def test_membru_sein_pozisaun_la_iha_alokasaun(self):
		self.jalan()
		self.assertFalse(Alokasaun.objects.filter(membro__nome='Hhh Sein Pozisaun').exists())
		self.assertFalse(Alokasaun.objects.filter(membro__nome='Ggg Membru Nasionál').exists())
		self.assertEqual(Membro.objects.filter(organizasaun=self.org).count(), 9)

	def test_formulariu_hatán_liu_telemovel(self):
		out = self.jalan()
		a = Membro.objects.get(nome='Aaa Presidente Uji')
		self.assertEqual(a.esperensia, 'Aaa pengalaman ONG')
		self.assertEqual(a.kontribuisaun, 'Hadia ekonomia')
		self.assertEqual(a.municipality_orijin, self.baucau)
		# naran iha aba Sheet1 ('Naran Seluk') la uza; liña ne'ebé la hatán la importa
		self.assertFalse(Membro.objects.filter(nome__startswith='Naran Seluk').exists())
		self.assertFalse(Membro.objects.filter(esperensia='Programador').exists())
		self.assertIn('la bele hatán', out)

	def test_postu_la_hetan_avizu(self):
		out = self.jalan()
		h = Membro.objects.get(nome='Hhh Sein Pozisaun')
		self.assertEqual(h.municipality, self.dili)
		self.assertIsNone(h.post)
		self.assertIn('Postu la hetan', out)

	def test_la_duplika_bainhira_la_o_dala_barak(self):
		self.jalan()
		n = (Membro.objects.count(), Alokasaun.objects.count(), Departamentu.objects.count())
		self.jalan()
		self.assertEqual((Membro.objects.count(), Alokasaun.objects.count(), Departamentu.objects.count()), n)

	def test_la_sobu_hakerek_dadus_edita(self):
		self.jalan()
		a = Membro.objects.get(nome='Aaa Presidente Uji')
		a.telefone = '70000000'
		a.save()
		self.jalan()
		a.refresh_from_db()
		self.assertEqual(a.telefone, '70000000')

	def test_dry_run_no_aprova(self):
		self.jalan('--dry-run')
		self.assertEqual(Membro.objects.count(), 0)
		self.assertEqual(Departamentu.objects.count(), 0)
		self.jalan('--aprova')
		self.assertTrue(Membro.objects.filter(aprovasaun='Aprovadu').count() == 9)

	def test_naran_menus_hanesan_hamutuk(self):
		linhas = list(LINHAS) + [(10, None, 'Aaa Presidente Ujii', None, 'Mane', None, None, 'Membru', 'Finansas', None, None, None, None, None, None, None, None, 'Nasional')]
		buat_xlsx(self.path, linhas)
		self.jalan()
		self.assertEqual(Membro.objects.count(), 9)
		self.assertTrue(Alokasaun.objects.filter(membro__nome='Aaa Presidente Uji', pozisaun='Membru').exists())

	def test_file_la_loos(self):
		with self.assertRaises(CommandError):
			call_command('import_see', '/la/iha.xlsx')
		from openpyxl import Workbook
		wb = Workbook()
		wb.save(self.path)
		with self.assertRaises(CommandError):
			self.jalan()

	def test_orgaun_la_iha(self):
		with self.assertRaises(CommandError):
			self.jalan('--orgaun', 'XYZ')
